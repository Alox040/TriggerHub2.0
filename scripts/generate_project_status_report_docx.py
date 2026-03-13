from __future__ import annotations

from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT_PATH = r"E:\Programmierung\TriggerHub2.0\docs\TriggerHub2_Projektstatusbericht_2026-03-11.docx"


def german_date(value: date) -> str:
    months = {
        1: "Januar",
        2: "Februar",
        3: "Maerz",
        4: "April",
        5: "Mai",
        6: "Juni",
        7: "Juli",
        8: "August",
        9: "September",
        10: "Oktober",
        11: "November",
        12: "Dezember",
    }
    return f"{value.day}. {months[value.month]} {value.year}"


def set_cell_bg(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_table_borders(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            tc = cell._tc
            tc_pr = tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{side}")
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "4")
                border.set(qn("w:space"), "0")
                border.set(qn("w:color"), "D9E2F0")
                borders.append(border)
            tc_pr.append(borders)


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if paragraph.runs:
        paragraph.runs[0].font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def add_paragraph(doc: Document, text: str, *, bold: bool = False, italic: bool = False) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(11)
    paragraph.paragraph_format.space_after = Pt(6)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        run.font.size = Pt(11)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    header_row = table.rows[0]
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        cell.text = header
        set_cell_bg(cell, "1F4E79")
        paragraph = cell.paragraphs[0]
        if paragraph.runs:
            paragraph.runs[0].font.bold = True
            paragraph.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            paragraph.runs[0].font.size = Pt(10)

    for row_index, row_data in enumerate(rows, start=1):
        fill = "EEF3F8" if row_index % 2 else "FFFFFF"
        row = table.rows[row_index]
        for cell_index, value in enumerate(row_data):
            cell = row.cells[cell_index]
            cell.text = value
            set_cell_bg(cell, fill)
            paragraph = cell.paragraphs[0]
            if paragraph.runs:
                paragraph.runs[0].font.size = Pt(10)

    set_table_borders(table)
    doc.add_paragraph()


doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)


# Title page
for _ in range(3):
    doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("TriggerHub 2.0")
run.font.size = Pt(28)
run.font.bold = True
run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Projektstatusbericht")
run.font.size = Pt(16)
run.font.italic = True
run.font.color.rgb = RGBColor(0x44, 0x72, 0xC4)

date_paragraph = doc.add_paragraph()
date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = date_paragraph.add_run(german_date(date.today()))
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_page_break()

# 2. Project overview
add_heading(doc, "Projektuebersicht", 1)
add_paragraph(
    doc,
    "TriggerHub 2.0 ist eine Software fuer Streamer und Content Creator. "
    "Sie soll wiederkehrende Arbeitsablaeufe vereinfachen und spaeter weitgehend automatisieren.",
)
add_paragraph(
    doc,
    "Gedacht ist das Projekt fuer Menschen, die beim Streamen oder Produzieren von Inhalten "
    "mehrere Programme gleichzeitig bedienen muessen, zum Beispiel fuer Aufnahme, Musik, Szenenwechsel "
    "oder spaetere Clip-Verarbeitung.",
)
add_paragraph(
    doc,
    "Das Problem dahinter ist einfach: Viele Creator verlieren Zeit und Konzentration durch manuelle "
    "Wiederholungsaufgaben. TriggerHub soll diese Ablaufe an einer Stelle zusammenfuehren.",
)

# 3. Goal
add_heading(doc, "Ziel des Projekts", 1)
add_paragraph(
    doc,
    "Die Vision von TriggerHub ist eine zentrale Schaltstelle fuer Creator-Workflows. "
    "Statt mehrere Einzeltools nebeneinander zu bedienen, soll der Nutzer spaeter wichtige Ablaufe "
    "einmal definieren und dann automatisiert ausfuehren lassen koennen.",
)
add_paragraph(
    doc,
    "Kurz gesagt: weniger manuelle Klicks, weniger Fehler unter Zeitdruck und mehr Fokus auf den eigentlichen Inhalt.",
)

# 4. Current development status
add_heading(doc, "Aktueller Entwicklungsstand", 1)
add_paragraph(
    doc,
    "Das Projekt ist aktuell in einer fortgeschrittenen Prototyp- und Vorproduktphase. "
    "Die technische Grundlage steht, viele Kernsysteme sind bereits vorhanden, aber einige Bereiche "
    "werden noch fuer den ersten breiteren Einsatz fertiggestellt.",
)
add_table(
    doc,
    ["Bereich", "Aktueller Stand", "Einordnung"],
    [
        [
            "Kernarchitektur",
            "Modulare Grundstruktur ist aufgebaut und verbindet alle Hauptmodule sauber miteinander.",
            "Stabil",
        ],
        [
            "Automationssystem",
            "Trigger-, Makro- und Event-Logik sind bereits vorhanden.",
            "Weit entwickelt",
        ],
        [
            "Benutzeroberflaeche",
            "Dashboard, Seitenstruktur und Grundkomponenten sind vorhanden, weitere Feinarbeit laeuft.",
            "In Ausbau",
        ],
        [
            "Desktop-App",
            "Electron-Huelle und Windows-Build-Pipeline sind eingerichtet.",
            "Technisch vorbereitet",
        ],
        [
            "Integrationen",
            "OBS-, Spotify- und Clip-Bausteine sind als technische Grundlage implementiert.",
            "Teilweise produktnah",
        ],
        [
            "Plugins",
            "Plugin-Registrierung und Erweiterungsstruktur sind angelegt.",
            "Grundlage vorhanden",
        ],
        [
            "Website",
            "Es gibt eine getrennte Website mit geschuetztem Owner-Only-Prelaunch-Zugang.",
            "Betriebsbereit fuer Prelaunch",
        ],
    ],
)

# 5. Finished
add_heading(doc, "Was bereits fertig ist", 1)
add_table(
    doc,
    ["Funktion oder Bereich", "Status"],
    [
        ["Grundarchitektur der Anwendung", "Implementiert"],
        ["Trigger-Engine fuer ereignisgesteuerte Ablaufe", "Implementiert"],
        ["Makro-System fuer mehrstufige Workflows", "Implementiert"],
        ["Zentrales Event-System", "Implementiert"],
        ["Plugin-Registry als Basis fuer Erweiterungen", "Implementiert"],
        ["OBS-, Spotify- und Clip-Service-Grundlagen", "Implementiert"],
        ["Desktop-Huelle fuer Windows", "Implementiert"],
        ["Release- und Build-Grundlagen", "Implementiert"],
        ["Getrennte Website mit Prelaunch-Schutz", "Implementiert"],
        ["Automatische Tests fuer zentrale Bereiche", "Implementiert"],
    ],
)
add_bullets(
    doc,
    [
        "Die Software hat bereits ein klares technisches Fundament.",
        "Die wichtigsten Kernbausteine existieren nicht nur als Idee, sondern als echter Code.",
        "Die Website ist vorhanden und fuer den aktuellen Owner-Only-Prelaunch abgesichert.",
    ],
)

# 6. Currently in development
add_heading(doc, "Was aktuell entwickelt wird", 1)
add_bullets(
    doc,
    [
        "Feinschliff fuer Release-Readiness der Desktop-App",
        "Weiterer Ausbau der Benutzeroberflaeche",
        "Praxisnahe Absicherung und Härtung der Integrationen",
        "Ausbau der Datenspeicherung und alltagsnahen Nutzbarkeit",
        "Weiterentwicklung der Website fuer spaetere Invite- und Public-Phasen",
    ],
)

# 7. Next steps
add_heading(doc, "Naechste Entwicklungsschritte", 1)
add_table(
    doc,
    ["Naechster Meilenstein", "Ziel"],
    [
        ["Desktop-Version weiter absichern", "Eine erste robuste Fassung fuer fruehe Tester bereitstellen."],
        ["Benutzeroberflaeche weiter vervollstaendigen", "Wichtige Arbeitsablaeufe einfacher und klarer bedienbar machen."],
        ["Reale Integrationen weiter pruefen", "Praxisnahe Nutzung mit echten Creator-Workflows verlaesslicher machen."],
        ["Datenspeicherung ausbauen", "Einstellungen und Ablaufe dauerhaft verfuegbar halten."],
        ["Website schrittweise erweitern", "Spaeter von owner-only zu invite_only und danach zu public_product uebergehen."],
    ],
)

# 8. Long-term vision
add_heading(doc, "Langfristige Vision", 1)
add_paragraph(
    doc,
    "Langfristig soll TriggerHub nicht nur eine einzelne App sein, sondern eine umfassende Plattform "
    "fuer Creator-Automation. Das Potenzial liegt darin, verschiedenste Tools, wiederkehrende Ablaufe "
    "und spaetere Erweiterungen an einer Stelle zusammenzufuehren.",
)
add_bullets(
    doc,
    [
        "Automationen fuer Streams und Content-Produktion zentral verwalten",
        "Mit Plugins und Erweiterungen wachsen koennen",
        "Wiederverwendbare Vorlagen und Workflows ermoeglichen",
        "Spaeter auch fuer mehr Nutzer als nur den Projekt-Owner geoeffnet werden",
    ],
)

# 9. Conclusion
add_heading(doc, "Fazit", 1)
add_paragraph(
    doc,
    "TriggerHub 2.0 ist bereits deutlich weiter als eine reine Idee. "
    "Die technische Basis steht und wichtige Kernsysteme wurden schon umgesetzt.",
)
add_paragraph(
    doc,
    "Besonders weit ist das Projekt bei der internen Struktur, den Automationsbausteinen "
    "und der allgemeinen Architektur.",
)
add_paragraph(
    doc,
    "Auch die Website ist nicht nur ein Entwurf, sondern besitzt bereits einen geschuetzten Prelaunch-Zugang.",
)
add_paragraph(
    doc,
    "Gleichzeitig befindet sich das Projekt noch nicht im Endzustand. "
    "Einige Bereiche werden gerade weiterentwickelt, damit der erste breitere Einsatz sinnvoll moeglich wird.",
)
add_paragraph(
    doc,
    "Die naechste Phase konzentriert sich darauf, die Anwendung robuster, nutzerfreundlicher und alltagstauglicher zu machen.",
)
add_paragraph(
    doc,
    "Danach kann das Projekt Schritt fuer Schritt von einem internen Owner-Only-Setup in einen groesseren Nutzungsrahmen wachsen.",
)
add_paragraph(
    doc,
    "Insgesamt ist der Entwicklungsstand solide, nachvollziehbar und klar in Richtung eines echten Produkts ausgerichtet.",
    bold=True,
)

doc.save(OUTPUT_PATH)
print(f"OK - DOCX gespeichert: {OUTPUT_PATH}")
